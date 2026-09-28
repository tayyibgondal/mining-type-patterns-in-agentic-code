"""Write reviewed labels into an annotation workbook.

Labels live in `datasets/evaluation/labels/<tag>_<stage>_labels.json` as
    {"<row number>": ["y"|"n", "reasoning"]}
so the judgements are version-controllable and reviewable on their own, rather
than being buried inside a binary xlsx.

The labels this writes into the Tayyib column were produced by reading each
pull request's title, body, parser construct counts and diff. They were made
WITHOUT looking at the pipeline's decision, which lives on the Predictions
sheet, so the evaluation is not circular. They are still a reviewed automatic
pass rather than an independent human annotation, and the paper must describe
them that way.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
EVAL = REPO / "datasets" / "evaluation"
LABELS = EVAL / "labels"

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def fill(tag: str, stage: str) -> None:
    book = EVAL / f"{tag}_{stage}_eval.xlsx"
    label_file = LABELS / f"{tag}_{stage}_labels.json"
    if not book.exists():
        raise SystemExit(f"missing workbook {book}")
    if not label_file.exists():
        raise SystemExit(f"missing labels {label_file}")

    labels = json.loads(label_file.read_text())
    wb = load_workbook(book)
    ev = wb["Evaluation"]
    header = [c.value for c in ev[1]]
    col_num = header.index("#") + 1
    col_ans = header.index("Tayyib (y/n)") + 1
    col_why = header.index("Tayyib_reasoning") + 1

    written = 0
    unlabelled = []
    for r in range(2, ev.max_row + 1):
        key = str(ev.cell(row=r, column=col_num).value)
        entry = labels.get(key)
        if not entry:
            unlabelled.append(key)
            continue
        answer, reason = entry
        ev.cell(row=r, column=col_ans, value=answer)
        ev.cell(row=r, column=col_why,
                value=_ILLEGAL.sub(" ", reason)[:32_000])
        written += 1

    wb.save(book)
    counts: dict[str, int] = {}
    for a, _ in labels.values():
        counts[a] = counts.get(a, 0) + 1
    print(f"[{tag}/{stage}] wrote {written} labels {counts} -> "
          f"{book.relative_to(REPO)}")
    if unlabelled:
        print(f"    still unlabelled: {unlabelled}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--stage", choices=["regex", "llm"], required=True)
    ap.add_argument("--tag", choices=["typescript", "csharp", "rust"],
                    action="append")
    args = ap.parse_args()
    for tag in (args.tag or ["typescript", "csharp", "rust"]):
        if (LABELS / f"{tag}_{args.stage}_labels.json").exists():
            fill(tag, args.stage)
        else:
            print(f"[{tag}/{args.stage}] no labels file yet, skipped")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
