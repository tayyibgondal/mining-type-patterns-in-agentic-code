"""Recompute the Metrics sheet of each annotation workbook, as real numbers.

Why numbers and not formulas: openpyxl writes formulas without a cached
result, so any viewer that does not recalculate on open shows empty cells.
Writing the computed values means the metrics are always visible.

The keep/drop decision is kept OUT of the workbook, in
`datasets/evaluation/labels/<tag>_<stage>_predictions.json`, so the Evaluation
sheet cannot anchor an annotator and the Metrics sheet needs no helper columns.

Run this after either annotator fills in or changes labels:

    python3 scripts/recompute_eval_metrics.py
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
EVAL = REPO / "datasets" / "evaluation"
LABELS = EVAL / "labels"
TAGS = {"typescript": "TypeScript", "csharp": "C#", "rust": "Rust"}
ANNOTATORS = ("Tayyib", "Imgyeong")
ROWS = ["TP", "FP", "TN", "FN", "Total", "Accuracy", "Precision", "Recall",
        "Specificity", "F1"]


def confusion(preds: dict[str, str], answers: dict[str, str]) -> dict:
    tp = fp = tn = fn = 0
    for key, pred in preds.items():
        ans = (answers.get(key) or "").strip().lower()
        if ans not in ("y", "n"):
            continue
        if pred == "keep" and ans == "y":
            tp += 1
        elif pred == "keep" and ans == "n":
            fp += 1
        elif pred == "drop" and ans == "n":
            tn += 1
        elif pred == "drop" and ans == "y":
            fn += 1
    total = tp + fp + tn + fn
    prec = tp / (tp + fp) if tp + fp else None
    rec = tp / (tp + fn) if tp + fn else None
    spec = tn / (tn + fp) if tn + fp else None
    f1 = (2 * prec * rec / (prec + rec)) if prec and rec else None
    return {"TP": tp, "FP": fp, "TN": tn, "FN": fn, "Total": total,
            "Accuracy": (tp + tn) / total if total else None,
            "Precision": prec, "Recall": rec, "Specificity": spec, "F1": f1}


def update(tag: str, stage: str, quiet: bool = False) -> dict | None:
    book = EVAL / f"{tag}_{stage}_eval.xlsx"
    pred_file = LABELS / f"{tag}_{stage}_predictions.json"
    if not book.exists() or not pred_file.exists():
        return None
    preds = json.loads(pred_file.read_text())

    wb = load_workbook(book)
    ev = wb["Evaluation"]
    header = [c.value for c in ev[1]]
    inum = header.index("#") + 1
    cols = {who: header.index(f"{who} (y/n)") + 1 for who in ANNOTATORS}

    answers = {who: {} for who in ANNOTATORS}
    for r in range(2, ev.max_row + 1):
        key = str(ev.cell(r, inum).value)
        for who in ANNOTATORS:
            answers[who][key] = ev.cell(r, cols[who]).value

    results = {who: confusion(preds, answers[who]) for who in ANNOTATORS}

    if "Metrics" in wb.sheetnames:
        del wb["Metrics"]
    mt = wb.create_sheet("Metrics")
    label = "regex parser" if stage == "regex" else "LLM filter"
    mt.cell(row=1, column=1,
            value=f"{label} performance - {TAGS[tag]}").font = Font(bold=True)
    mt.cell(row=2, column=1, value=(
        "Positive class = PR genuinely involves construct work. Values are "
        "computed; re-run scripts/recompute_eval_metrics.py after annotating."))
    mt.cell(row=4, column=1, value="Metric").font = Font(bold=True)
    for col, who in enumerate(ANNOTATORS, start=2):
        mt.cell(row=4, column=col, value=who).font = Font(bold=True)
    for i, name in enumerate(ROWS, start=5):
        mt.cell(row=i, column=1, value=name)
        for col, who in enumerate(ANNOTATORS, start=2):
            value = results[who][name]
            cell = mt.cell(row=i, column=col,
                           value="" if value is None else value)
            if name in ("Accuracy", "Precision", "Recall", "Specificity", "F1"):
                cell.number_format = "0.0%"

    agree = both = 0
    for key in preds:
        a = (answers["Tayyib"].get(key) or "").strip().lower()
        b = (answers["Imgyeong"].get(key) or "").strip().lower()
        if a in ("y", "n") and b in ("y", "n"):
            both += 1
            agree += int(a == b)
    mt.cell(row=16, column=1, value="Raw agreement (rows both filled)")
    cell = mt.cell(row=16, column=2, value=(agree / both) if both else "")
    cell.number_format = "0.0%"
    mt.cell(row=17, column=1, value="Rows both filled")
    mt.cell(row=17, column=2, value=both)

    mt.column_dimensions["A"].width = 52
    for col in range(2, 2 + len(ANNOTATORS)):
        mt.column_dimensions[get_column_letter(col)].width = 15
    wb.save(book)

    if not quiet:
        t = results["Tayyib"]
        line = (f"  {tag}/{stage:<6} TP={t['TP']:>3} FP={t['FP']:>3} "
                f"TN={t['TN']:>3} FN={t['FN']:>3}")
        for name in ("Accuracy", "Precision", "Recall", "F1"):
            v = t[name]
            line += f"  {name[:4]}={'--' if v is None else f'{v:.1%}'}"
        print(line)
    return results


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--stage", choices=["regex", "llm"], action="append")
    ap.add_argument("--tag", choices=list(TAGS), action="append")
    args = ap.parse_args()
    print("Tayyib column:")
    for stage in (args.stage or ["regex", "llm"]):
        for tag in (args.tag or list(TAGS)):
            update(tag, stage)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
